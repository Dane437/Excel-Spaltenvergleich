from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
import unicodedata
import pandas as pd
import re
from datetime import datetime, time

# Vergleicht die Werte in den angegebenen Spalten und hebt die Unterschiede hervor
def highlight_differences(file_name, sheet_name, col_left, col_right, header_row, valid_combinations, empty_is_difference):
    wb = load_workbook(file_name)
    ws = wb[sheet_name]

    headers = {cell.value: cell.column for cell in ws[header_row]}
    rows_different = []

    for row in range(header_row + 1, ws.max_row + 1):
        c1 = ws.cell(row=row, column=headers[col_left])
        c2 = ws.cell(row=row, column=headers[col_right])
        
        if not empty_is_difference and (c1.value is None or c2.value is None):
            #continue
            one_is_empty = True

        if not valid_combinations:
            if c1.value != c2.value:
                rows_different.append(row)
                if one_is_empty:
                    c1.fill = PatternFill(start_color='FFEB9C', fill_type='solid')  # gelbe Formatierung
                    c2.fill = PatternFill(start_color='FFEB9C', fill_type='solid')  # gelbe Formatierung
                    one_is_empty = False
                else:
                    c1.fill = PatternFill(start_color='FFCCCC', fill_type='solid')  # rote Formatierung
                    c2.fill = PatternFill(start_color='FFCCCC', fill_type='solid')  # rote Formatierung
        else:
            if (c1.value, c2.value) in valid_combinations:
                continue

            c1.fill = PatternFill(start_color='FFCCCC', fill_type='solid')  # rote Formatierung
            c2.fill = PatternFill(start_color='FFCCCC', fill_type='solid')  # rote Formatierung

    wb.save(file_name)
    return rows_different


# Hilfsfunktion zum Sortieren von Strings mit deutschen Umlauten
def de_sort_key(s):
    return (
        unicodedata.normalize("NFKD", s)
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )


def create_differences_sheet(file_name, rows_different, header_row):
    wb = load_workbook(file_name)
    ws = wb['Temp']
    headers = {cell.value: cell.column for cell in ws[header_row]}
    worksheet_copy = wb.copy_worksheet(ws)
    worksheet_copy.title = 'Unterschiede'

    for row in range(worksheet_copy.max_row, header_row, -1):    # zählt rückwärts und stoppt bei header_row
        if row not in rows_different:
             worksheet_copy.delete_rows(row)

    wb.save(file_name)

def bereinige_hausnummer(x):
    if pd.isna(x):
        return x

    # Falls Excel als Uhrzeit erkannt hat
    if isinstance(x, time):
        return f"{x.hour}a"

    # Falls Excel als Datum erkannt hat
    if isinstance(x, datetime):
        return f"{x.day}-{x.month}"

    s = str(x).strip()

   # Alle Leerzeichen entfernen
    s = re.sub(r"\s+", "", s)

    # Alles in Kleinbuchstaben umwandeln
    s = s.lower()

    # Uhrzeit als Text: 01:00:00 bis 23:00:00 -> 1a bis 23a
    m = re.fullmatch(r'(\d{1,2}):00:00', s)
    if m:
        stunde = int(m.group(1))
        return f"{stunde}a"

    # Datum als Text: 01.01.2026, 31.12.2026 usw.
    m = re.fullmatch(r'(\d{1,2})\.(\d{1,2})\.\d{4}', s)
    if m:
        tag = int(m.group(1))
        monat = int(m.group(2))
        return f"{tag}-{monat}"

    # Falls das Jahr fehlt: 01.01.
    m = re.fullmatch(r'(\d{1,2})\.(\d{1,2})\.?', s)
    if m:
        tag = int(m.group(1))
        monat = int(m.group(2))
        return f"{tag}-{monat}"

    return s